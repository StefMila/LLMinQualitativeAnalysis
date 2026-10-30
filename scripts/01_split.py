"""Passo 1: dall'Excel a serbatoio / valutazione / resto, con seed fisso.

Lavora sulla domanda indicata da dati.domanda_attiva in config.yaml.
Output (colonne id, testo, umano): data/splits/<domanda>/serbatoio.csv, valutazione.csv, resto.csv
"""
import pandas as pd
from ftfy import fix_text
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from common import DOMANDA, RESULTS, ROOT, SPLITS, carica_config, tutte_le_etichette

cfg = carica_config()
d = cfg["dati"]
q = d["domande"][DOMANDA]

percorso_excel = ROOT / d["excel"]
df = pd.read_excel(percorso_excel, sheet_name=d["foglio"])
print(f"Domanda: {DOMANDA} ({q['testo_domanda']})")
print(f"Righe nell'Excel: {len(df)}")

# Testi con codifica corrotta (es. "Ã¨" al posto di "è"): si correggono con ftfy. L'Excel originale
# resta intatto; le correzioni vanno in una copia *_ftfy.xlsx e in correzioni_mojibake.csv.
MOJIBAKE = ("Ã", "â€", "Â")
correzioni, residui = [], []
for nome_domanda, dettagli in d["domande"].items():
    colonna = column_index_from_string(dettagli["col_testo"]) - 1
    for riga, valore in enumerate(df.iloc[:, colonna]):
        if pd.isna(valore):
            continue
        prima = str(valore)
        dopo = fix_text(prima) if any(s in prima for s in MOJIBAKE) else prima
        if dopo != prima:
            df.iat[riga, colonna] = dopo
            correzioni.append({"id": riga + 2, "domanda": nome_domanda, "testo_prima": prima, "testo_dopo": dopo})
        if any(s in dopo for s in MOJIBAKE):
            residui.append(f"ID {riga + 2} ({nome_domanda})")
if residui:
    raise ValueError(f"Mojibake ancora presente dopo ftfy: {', '.join(residui[:10])}")
if correzioni:
    workbook = load_workbook(percorso_excel)
    foglio = workbook[d["foglio"]]
    for c in correzioni:
        colonna = column_index_from_string(d["domande"][c["domanda"]]["col_testo"])
        foglio.cell(row=c["id"], column=colonna).value = c["testo_dopo"]
    percorso_corretto = percorso_excel.with_name(f"{percorso_excel.stem}_ftfy{percorso_excel.suffix}")
    workbook.save(percorso_corretto)
    RESULTS.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(correzioni).to_csv(RESULTS / "correzioni_mojibake.csv", index=False, encoding="utf-8-sig")
    print(f"Dataset originale intatto; copia corretta: {percorso_corretto}")
    print(f"Log correzioni: {RESULTS / 'correzioni_mojibake.csv'}")
else:
    print("Controllo mojibake: nessuna sequenza sospetta dopo ftfy.")


def colonna(lettera):
    """La colonna dell'Excel indicata con la lettera (es. "BD")."""
    return df.iloc[:, column_index_from_string(lettera) - 1]


df = pd.DataFrame({
    "id": df.index + 2,  # numero di riga in Excel (riga 1 = intestazioni): così si ritrova la risposta nel file
    "testo": colonna(q["col_testo"]),
    "umano": colonna(q["col_umano"]),
})

# Risposte vuote: non hanno senso né per il modello né per il confronto
vuote = df["testo"].isna() | (df["testo"].astype(str).str.strip() == "")
if vuote.any():
    print(f"Scarto {vuote.sum()} risposte vuote.")
df = df[~vuote]

# Coding 1 = "-": risposta non ancora codificata, senza riferimento umano
non_codificate = df["umano"].astype(str).str.strip() == d["non_codificato"]
if non_codificate.any():
    print(f"Scarto {non_codificate.sum()} risposte non ancora codificate ('{d['non_codificato']}').")
df = df[~non_codificate]

# Coding 1 vuoto = non classificabile (vuoti, non so, rumore)
df["umano"] = df["umano"].fillna(cfg["non_classificabile"]).astype(str).str.strip()

ignote = set(df["umano"]) - set(tutte_le_etichette(cfg))
if ignote:
    print(f"ATTENZIONE: etichette umane non presenti in config.yaml: {sorted(ignote)}")

print(f"Risposte utilizzabili: {len(df)}")
print("Distribuzione della codifica umana:")
print(df["umano"].value_counts().to_string())

# Mescola con seed fisso e taglia
df = df.sample(frac=1, random_state=cfg["seed"]).reset_index(drop=True)
n_s, n_v = cfg["split"]["n_serbatoio"], cfg["split"]["n_valutazione"]
assert n_s + n_v <= len(df), f"Servono almeno {n_s + n_v} risposte, ne ho {len(df)}."

SPLITS.mkdir(parents=True, exist_ok=True)
for nome, parte in [("serbatoio", df.iloc[:n_s]), ("valutazione", df.iloc[n_s:n_s + n_v]), ("resto", df.iloc[n_s + n_v:])]:
    parte.to_csv(SPLITS / f"{nome}.csv", index=False, encoding="utf-8-sig")
    print(f"{nome:12s} {len(parte):5d} risposte")
